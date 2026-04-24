from __future__ import annotations

from io import StringIO, BytesIO
from pathlib import Path
from typing import Any, Dict, List, Tuple, Optional

import pandas as pd
from PyPDF2 import PdfReader
from docx import Document

from src.classifier import classify_document
from src.summarizer import summarize_text
from src.keyword_extractor import extract_keywords
from src.risk_engine import calculate_risk
from src.text_cleaner import clean_extracted_text

try:
    import pytesseract
except Exception:
    pytesseract = None

try:
    from pdf2image import convert_from_bytes
except Exception:
    convert_from_bytes = None


SUPPORTED_EXTENSIONS = {"txt", "csv", "pdf", "docx", "xlsx"}


class OCRUnavailableError(RuntimeError):
    pass


def get_file_extension(filename: str) -> str:
    return Path(filename).suffix.lower().lstrip(".")


def safe_decode(raw: bytes) -> str:
    for encoding in ["utf-8", "utf-8-sig", "cp1252", "latin-1"]:
        try:
            return raw.decode(encoding, errors="ignore").strip()
        except Exception:
            continue
    return ""


def extract_text_from_txt_bytes(data: bytes) -> str:
    return safe_decode(data)


def extract_text_from_csv_bytes(data: bytes) -> str:
    raw_text = safe_decode(data)
    if not raw_text:
        return "CSV file is empty."
    try:
        df = pd.read_csv(StringIO(raw_text))
    except Exception:
        try:
            df = pd.read_csv(StringIO(raw_text), sep=";")
        except Exception:
            return raw_text

    if df.empty:
        return "CSV file is empty."

    lines = []
    for _, row in df.fillna("").iterrows():
        row_text = " | ".join(f"{col}: {row[col]}" for col in df.columns)
        lines.append(row_text)
    return "\n".join(lines).strip()


def extract_text_from_xlsx_bytes(data: bytes) -> str:
    try:
        sheets = pd.read_excel(BytesIO(data), sheet_name=None, dtype=str)
    except Exception as exc:
        return f"Excel file could not be read: {exc}"
    if not sheets:
        return "Excel file is empty."
    parts = []
    for sheet_name, df in sheets.items():
        df = df.fillna("")
        if df.empty:
            continue
        lines = [f"Sheet: {sheet_name}"]
        for _, row in df.iterrows():
            row_text = " | ".join(f"{col}: {row[col]}" for col in df.columns)
            lines.append(row_text)
        parts.append("\n".join(lines))
    return "\n\n".join(parts).strip() or "Excel file is empty."


def ocr_status() -> dict:
    return {
        "pytesseract": pytesseract is not None,
        "pdf2image": convert_from_bytes is not None,
        "available": pytesseract is not None and convert_from_bytes is not None,
    }


def ocr_pdf_bytes(data: bytes, *, dpi: int = 220, page_limit: int = 30) -> Tuple[str, bool]:
    if pytesseract is None or convert_from_bytes is None:
        raise OCRUnavailableError(
            "OCR dependencies are not available. Install pytesseract and pdf2image, plus native Tesseract OCR and Poppler."
        )
    images = convert_from_bytes(data, dpi=dpi)
    texts = []
    for img in images[: max(1, page_limit)]:
        texts.append(pytesseract.image_to_string(img))
    merged = "\n".join(t.strip() for t in texts if t and t.strip()).strip()
    return merged, bool(merged)


def extract_text_from_pdf_bytes(
    data: bytes,
    *,
    ocr_enabled: bool = True,
    ocr_page_limit: int = 30,
    ocr_dpi: int = 220,
    min_words_threshold: int = 30,
) -> Tuple[str, bool]:
    reader = PdfReader(BytesIO(data))
    pages_text = []
    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        if page_text.strip():
            pages_text.append(page_text.strip())

    merged = "\n\n".join(pages_text).strip()
    if len(merged.split()) >= max(1, min_words_threshold):
        return merged, False

    if not ocr_enabled:
        return merged or "No readable text found in PDF.", False

    try:
        ocr_text, ok = ocr_pdf_bytes(data, dpi=ocr_dpi, page_limit=ocr_page_limit)
        if ok:
            return ocr_text, True
    except OCRUnavailableError:
        pass
    except Exception:
        pass
    return merged or "No readable text found in PDF.", False


def extract_text_from_docx_bytes(data: bytes) -> str:
    doc = Document(BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    table_lines = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                table_lines.append(" | ".join(cells))

    content_parts = []
    if paragraphs:
        content_parts.append("\n".join(paragraphs))
    if table_lines:
        content_parts.append("\n".join(table_lines))
    return "\n\n".join(content_parts).strip() or "No readable text found in DOCX."


def normalize_analysis_text(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\x00", " ").replace("\r", " ").replace("\t", " ")
    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    return text.strip()


def build_management_takeaway(document_type: str, risk_label: str, risk_categories: List[str], keywords: List[str]) -> str:
    category_text = ", ".join(risk_categories[:3]) if risk_categories else "no strong risk categories"
    keyword_text = ", ".join(keywords[:5]) if keywords else "limited dominant keywords"
    return (
        f"This file is classified as {document_type} with {risk_label} risk. "
        f"Primary risk areas include {category_text}. "
        f"Key terms detected include {keyword_text}."
    )


def build_risk_explanation(risk_label: str, risk_score: int, risk_categories: List[str], flagged_phrases: List[str]) -> str:
    cats = ", ".join(risk_categories[:4]) if risk_categories else "No strong categories detected"
    flags = ", ".join(flagged_phrases[:8]) if flagged_phrases else "no major triggered phrases"
    return f"Primary categories: {cats}. Triggered phrases: {flags}. Assigned label {risk_label} with score {risk_score}."


def apply_mode_processing(analysis_text: str, mode: str) -> dict:
    normalized_mode = (mode or "Premium").strip().title()
    if normalized_mode not in {"Lite", "Premium"}:
        normalized_mode = "Premium"

    document_type = classify_document(analysis_text)
    summary = summarize_text(analysis_text, max_sentences=4 if normalized_mode == "Premium" else 2)
    keywords = extract_keywords(analysis_text, top_n=10 if normalized_mode == "Premium" else 5)
    full_risk = calculate_risk(analysis_text)
    if normalized_mode == "Premium":
        risk_result = full_risk
    else:
        risk_result = {
            "risk_score": full_risk["risk_score"],
            "risk_label": full_risk["risk_label"],
            "risk_categories": full_risk["risk_categories"][:3],
            "flagged_phrases": full_risk["flagged_phrases"][:5],
        }

    management_takeaway = build_management_takeaway(
        document_type=document_type,
        risk_label=risk_result["risk_label"],
        risk_categories=risk_result["risk_categories"],
        keywords=keywords,
    )
    risk_explanation = build_risk_explanation(
        risk_label=risk_result["risk_label"],
        risk_score=int(risk_result["risk_score"]),
        risk_categories=risk_result["risk_categories"],
        flagged_phrases=risk_result["flagged_phrases"],
    )

    return {
        "document_type": document_type,
        "summary": summary,
        "keywords": keywords,
        "risk_score": risk_result["risk_score"],
        "risk_label": risk_result["risk_label"],
        "risk_categories": risk_result["risk_categories"],
        "flagged_phrases": risk_result["flagged_phrases"],
        "management_takeaway": management_takeaway,
        "risk_explanation": risk_explanation,
        "mode": normalized_mode,
    }


def process_file_bytes(
    filename: str,
    data: bytes,
    mode: str = "Premium",
    ocr_config: Optional[dict] = None,
    file_path: Optional[str] = None,
) -> Dict[str, Any]:
    extension = get_file_extension(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: .{extension}")

    ocr_used = False
    ocr_config = ocr_config or {}
    if extension == "txt":
        raw_content = extract_text_from_txt_bytes(data)
    elif extension == "csv":
        raw_content = extract_text_from_csv_bytes(data)
    elif extension == "pdf":
        raw_content, ocr_used = extract_text_from_pdf_bytes(
            data,
            ocr_enabled=bool(ocr_config.get("ocr_enabled", True)),
            ocr_page_limit=int(ocr_config.get("ocr_page_limit", 30)),
            ocr_dpi=int(ocr_config.get("ocr_dpi", 220)),
            min_words_threshold=int(ocr_config.get("ocr_min_words_threshold", 30)),
        )
    elif extension == "docx":
        raw_content = extract_text_from_docx_bytes(data)
    elif extension == "xlsx":
        raw_content = extract_text_from_xlsx_bytes(data)
    else:
        raise ValueError(f"Unsupported file type: .{extension}")

    raw_content = normalize_analysis_text(raw_content)
    cleaned_content = clean_extracted_text(raw_content)
    cleaned_content = normalize_analysis_text(cleaned_content)
    analysis_text = cleaned_content if cleaned_content else raw_content
    if not analysis_text:
        analysis_text = "No readable text content could be extracted from this file."

    mode_result = apply_mode_processing(analysis_text, mode)

    return {
        "filename": filename,
        "file_path": file_path or f"uploaded://{filename}",
        "filetype": extension.upper(),
        "mode": mode_result["mode"],
        "content": analysis_text,
        "raw_content": raw_content,
        "cleaned_content": cleaned_content,
        "char_count": len(analysis_text),
        "word_count": len(analysis_text.split()) if analysis_text else 0,
        "document_type": mode_result["document_type"],
        "summary": mode_result["summary"],
        "keywords": mode_result["keywords"],
        "risk_score": mode_result["risk_score"],
        "risk_label": mode_result["risk_label"],
        "risk_categories": mode_result["risk_categories"],
        "flagged_phrases": mode_result["flagged_phrases"],
        "management_takeaway": mode_result["management_takeaway"],
        "risk_explanation": mode_result["risk_explanation"],
        "ocr_used": ocr_used,
    }


def process_file_path(file_path: str, mode: str = "Premium", ocr_config: Optional[dict] = None) -> Dict[str, Any]:
    path = Path(file_path)
    data = path.read_bytes()
    return process_file_bytes(path.name, data, mode=mode, ocr_config=ocr_config, file_path=file_path)


def process_uploaded_file(uploaded_file, mode: str = "Premium", ocr_config: Optional[dict] = None) -> Dict[str, Any]:
    return process_file_bytes(
        uploaded_file.name,
        uploaded_file.getvalue(),
        mode=mode,
        ocr_config=ocr_config,
        file_path=f"uploaded://{uploaded_file.name}",
    )


def process_multiple_files(uploaded_files, mode: str = "Premium", ocr_config: Optional[dict] = None) -> List[Dict[str, Any]]:
    return [process_uploaded_file(f, mode=mode, ocr_config=ocr_config) for f in uploaded_files]
